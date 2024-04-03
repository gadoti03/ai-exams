import os
from typing import Dict, List, Any

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm
from docx.styles.style import ParagraphStyle

from src.csp import CSP
from src.exercise import Exercise, Dummy
from src.game import Game
from src.planning import Planning
from src.questions import Questions
from src.search import Search
from src.training import Training


class Exam:

    def __init__(self, sources: str, exports: str, exam: str):
        """A class representing the whole exam to be exported."""

        # load the configuration
        path = os.path.join(sources, exam)
        with open(path, 'r') as file:
            config = yaml.safe_load(file)

        self.config: Dict[str, Any] = config
        """The configuration file of the exam."""

        self.path: str = os.path.join(exports, exam.replace('.yml', '.docx'))
        """The filepath of the output document."""

        self.document: Document = Document()
        """The word document representing the exam."""

    @property
    def title(self) -> ParagraphStyle:
        """The style of the main title."""
        return self.document.styles['Heading 1']

    @property
    def exercise(self) -> ParagraphStyle:
        """The style of the exercise titles."""
        return self.document.styles['Heading 2']

    @property
    def text(self) -> ParagraphStyle:
        """The style of the exercise text."""
        return self.document.styles['Normal']

    def save(self):
        """Processes and exports the final exam in docx format."""
        # parse exercises from yaml configuration
        exercises = self._parse_exercises()
        # write exam text
        title = f"EXAM OF {self.config['exam']}\n{self.config['date']}\nPROF. MICHELA MILANO"
        self.document.add_paragraph(title, style=self.title)
        for i, x in enumerate(exercises):
            self.document.add_paragraph(f'Exercise {i + 1}', style=self.exercise)
            x.text(self.document)
        # write exam solution
        self.document.add_page_break()
        self.document.add_paragraph("SOLUTION", style=self.title)
        for i, x in enumerate(exercises):
            self.document.add_paragraph(f'Exercise {i + 1}', style=self.exercise)
            x.solution(self.document)
            if i + 1 != len(exercises):
                self.document.add_page_break()
        # format text and save document
        self._format_and_save()

    def _parse_exercises(self) -> List[Exercise]:
        """Parses the exercises from the yaml file."""
        output = []
        planning = None  # keep a reference to the planning exercise to be passed to the questions exercise
        for exercise, kwargs in self.config['exercises'].items():
            if exercise == 'dummy':
                assert kwargs is None, "No arguments expected for empty exercise"
                exercise = Dummy()
            elif exercise == 'game':
                assert kwargs is None, "No arguments expected for game exercise"
                exercise = Game()
            elif exercise == 'search':
                exercise = Search(**kwargs)
            elif exercise == 'csp':
                exercise = CSP(**kwargs)
            elif exercise == 'planning':
                exercise = Planning(**kwargs)
                planning = exercise
            elif exercise == 'questions':
                exercise = Questions(planning=planning, **kwargs)
            elif exercise == 'training':
                exercise = Training(**kwargs)
            else:
                raise AssertionError(f"Unknown exercise '{exercise}'")
            output.append(exercise)
        return output

    def _format_and_save(self):
        """Sets the formatting of paragraphs and saves the document."""
        # TITLE
        style = self.title
        font = style.font
        font.size = Pt(14)
        font.name = 'Calibri'
        font.color.rgb = RGBColor(0, 0, 0)
        font.bold = True
        form = style.paragraph_format
        form.alignment = WD_ALIGN_PARAGRAPH.CENTER
        form.space_after = 0
        form.space_before = 0
        form.line_spacing = 1.0
        # EXERCISE
        style = self.exercise
        font = style.font
        font.size = Pt(12)
        font.name = 'Calibri'
        font.color.rgb = RGBColor(0, 0, 0)
        font.bold = True
        form = style.paragraph_format
        form.alignment = WD_ALIGN_PARAGRAPH.LEFT
        form.space_after = Pt(12)
        form.space_before = Pt(12)
        form.line_spacing = 1.0
        form.keep_with_next = True  # keep this paragraph in the same page as next paragraph (text)
        # TEXT
        style = self.text
        font = style.font
        font.size = Pt(12)
        font.name = 'Times New Roman'
        font.color.rgb = RGBColor(0, 0, 0)
        font.bold = False
        form = style.paragraph_format
        form.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        form.space_after = 0
        form.space_before = 0
        form.line_spacing = 1.0
        # SET DIMENSIONS AND MARGINS
        sections = self.document.sections
        for section in sections:
            section.page_height = Cm(29.7)
            section.page_width = Cm(21.0)
            section.top_margin = Cm(1.0)
            section.bottom_margin = Cm(1.27)
            section.left_margin = Cm(1.27)
            section.right_margin = Cm(1.27)
        # EXPORT
        self.document.save(self.path)
